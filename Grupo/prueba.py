from openpyxl import Workbook
from openpyxl import load_workbook

# libro2 = load_workbook("libro.xlsx")
# hoja = libro2.active

# for i in hoja.iter_rows(values_only=True):
#     print(i)
#     # hoja.append(i)
# hoja.cell(row=1, column=1, value="Holaaaa")
# print(hoja.cell(row=1, column=1).value)

# libro2.save("libro2.xlsx")
# print("Archivo mandado")
    
    
'''CRUD
C - create Workbook
R - read load_workbook
U - update load_workbook - crear uno nuevo(modificado) - eliminar(antiguo)
D - delete reemplazas la celda por un null, None, o ""
'''


# libro = Workbook()
# hoja = libro.active
# hoja.title = "Hoja1"

# hoja.append([1,2,3,4,5,6,7,8,9])
# hoja.append([x for x in range(10)])
# hoja.append([x for x in range(10)])
# hoja.append([x for x in range(10)])
# hoja.append([x for x in range(10)])
# hoja.append([x for x in range(10)])
# hoja.append([x for x in range(10)])
# hoja.append([x for x in range(10)])
# hoja.append([x for x in range(10)])
# hoja.append([x for x in range(10)])
# hoja.append([x for x in range(10)])

# libro.save("libro.xlsx")
# print("Listo, se mandó")
